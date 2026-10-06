#!/usr/bin/env python3
"""Build the meetup deck. Facts retrieved 6 October 2026. See ../sources.md.

    python3 -m pip install python-pptx
    python3 deck/build_deck.py
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Calibri"
W = 13.333
H = 7.5

PAPER = RGBColor(0xF6, 0xF3, 0xEC)
INK = RGBColor(0x1C, 0x19, 0x15)
MUTED = RGBColor(0x6B, 0x64, 0x5C)
ACCENT = RGBColor(0x0E, 0x4B, 0x48)
ACCENT_SOFT = RGBColor(0xE4, 0xF0, 0xEE)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
CARD = RGBColor(0xFF, 0xFC, 0xF7)
LINE = RGBColor(0xE0, 0xD9, 0xCC)
BAND = RGBColor(0xEF, 0xEA, 0xE1)
LABEL = RGBColor(0xE7, 0xE1, 0xD4)
HEADER_INK = RGBColor(0xF6, 0xF3, 0xEC)

OUT = Path(__file__).resolve().parent / "ai-driven-development.pptx"
TOTAL = 20


def style_run(run, text, size, bold, italic, color):
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = r_pr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(r_pr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def set_anchor(tf, anchor: str):
    tf._txBody.bodyPr.set("anchor", anchor)
    tf.word_wrap = True
    tf.auto_size = None


def fill_tf(tf, paragraphs, align=PP_ALIGN.LEFT, anchor="t"):
    """paragraphs: list of dicts with text, size, bold, italic, color, before, after."""
    set_anchor(tf, anchor)
    for index, block in enumerate(paragraphs):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = block.get("align", align)
        p.space_before = Pt(block.get("before", 0))
        p.space_after = Pt(block.get("after", 0))
        if "spacing" in block:
            p.line_spacing = block["spacing"]
        run = p.add_run()
        style_run(
            run,
            block["text"],
            block.get("size", 18),
            block.get("bold", False),
            block.get("italic", False),
            block.get("color", INK),
        )


def add_text(slide, paragraphs, x, y, w, h, align=PP_ALIGN.LEFT, anchor="t"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    if isinstance(paragraphs, str):
        paragraphs = [{"text": paragraphs}]
    fill_tf(tf, paragraphs, align=align, anchor=anchor)
    return shape


def paint(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill, line=None, radius=None, shape=MSO_SHAPE.RECTANGLE):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
        item.line.width = Pt(1)
    if radius is not None:
        try:
            item.adjustments[0] = radius
        except Exception:
            pass
    return item


def shape_text(shape, paragraphs, align=PP_ALIGN.LEFT, anchor="ctr", size=14):
    tf = shape.text_frame
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    if isinstance(paragraphs, str):
        paragraphs = [{"text": paragraphs, "size": size}]
    fill_tf(tf, paragraphs, align=align, anchor=anchor)


def bar(slide):
    rect(slide, 0, 0, W, 0.08, ACCENT)


def footer(slide, number, label="AI-driven development"):
    add_text(
        slide,
        [{"text": label, "size": 12, "color": MUTED}],
        0.62,
        7.12,
        8.5,
        0.26,
        anchor="ctr",
    )
    add_text(
        slide,
        [{"text": f"{number}  /  {TOTAL}", "size": 12, "color": MUTED}],
        9.4,
        7.12,
        3.3,
        0.26,
        align=PP_ALIGN.RIGHT,
        anchor="ctr",
    )


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def kicker_title(slide, kicker, title):
    add_text(
        slide,
        [{"text": kicker.upper(), "size": 13, "bold": True, "color": ACCENT}],
        0.62,
        0.28,
        12.1,
        0.28,
    )
    add_text(
        slide,
        [{"text": title, "size": 32, "bold": True}],
        0.62,
        0.54,
        12.1,
        0.62,
    )


def flow(slide, labels, x, y, w, h=0.48, size=14):
    count = len(labels)
    gap = 0.32
    pill_w = (w - gap * (count - 1)) / count
    for index, label in enumerate(labels):
        px = x + index * (pill_w + gap)
        last = index == count - 1
        pill = rect(
            slide,
            px,
            y,
            pill_w,
            h,
            ACCENT if last else ACCENT_SOFT,
            radius=0.5,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        )
        shape_text(
            pill,
            [
                {
                    "text": label,
                    "size": size,
                    "bold": True,
                    "color": HEADER_INK if last else INK,
                }
            ],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        if index < count - 1:
            add_text(
                slide,
                [{"text": "→", "size": 16, "bold": True, "color": ACCENT}],
                px + pill_w,
                y,
                gap,
                h,
                align=PP_ALIGN.CENTER,
                anchor="ctr",
            )


def bullets(slide, items, x, y, w, h, size=18):
    paragraphs = []
    for index, item in enumerate(items):
        paragraphs.append(
            {
                "text": item,
                "size": size,
                "before": 0 if index == 0 else 12,
                "after": 0,
            }
        )
    add_text(slide, paragraphs, x, y, w, h)


def meta_line(slide, text):
    add_text(
        slide,
        [{"text": text, "size": 14, "color": MUTED}],
        0.62,
        6.62,
        12.1,
        0.36,
        anchor="ctr",
    )


def new_slide(prs, number, spoken):
    layout = next(l for l in prs.slide_layouts if l.name == "Blank")
    slide = prs.slides.add_slide(layout)
    paint(slide)
    bar(slide)
    footer(slide, number)
    notes(slide, spoken)
    return slide


def slide_title(prs):
    slide = new_slide(
        prs,
        1,
        """
        Open with the question, not the answer. Name the four things people are adopting: Spec Kit, BMAD, Squad, and a three-file loop. Say this will take about half an hour, and that the appendix has every figure if someone wants the source. Do not preview the verdict yet.
        """,
    )
    add_text(
        slide,
        [{"text": "SOFTWARE ENGINEERING MEETUP", "size": 14, "bold": True, "color": ACCENT}],
        0.7,
        1.35,
        12,
        0.32,
    )
    add_text(
        slide,
        [{"text": "Do the new AI frameworks\nearn their keep?", "size": 44, "bold": True}],
        0.7,
        1.8,
        12,
        1.7,
    )
    add_text(
        slide,
        [
            {
                "text": "Spec Kit, BMAD, Squad, and three files you keep current.",
                "size": 22,
                "color": MUTED,
            }
        ],
        0.7,
        3.7,
        11.5,
        0.5,
    )
    flow(
        slide,
        ["Requirement", "Short plan", "Code", "Update the plan"],
        0.7,
        4.7,
        11.9,
        h=0.56,
        size=16,
    )
    add_text(
        slide,
        [
            {
                "text": "Planning is worth more. The artifact should be shorter.",
                "size": 16,
                "color": MUTED,
            }
        ],
        0.7,
        5.5,
        11.5,
        0.36,
    )


def slide_shift(prs):
    slide = new_slide(
        prs,
        2,
        """
        The model will build whatever assumption you left unstated, and it will look finished. That is why planning is worth more than it was when typing was the bottleneck. The shape of the plan changed with it. A long document written once goes stale and the agent still treats it as truth. A short plan you update when the work changes is the planning that got more valuable.
        """,
    )
    kicker_title(slide, "The shift", "Intent got expensive")
    left = rect(slide, 0.62, 1.7, 5.85, 4.5, CARD, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    right = rect(slide, 6.85, 1.7, 5.85, 4.5, ACCENT, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # redraw text on top via text frames of the cards
    shape_text(
        left,
        [
            {"text": "Written once", "size": 14, "bold": True, "color": ACCENT, "before": 0, "after": 8},
            {"text": "A long specification", "size": 28, "bold": True, "before": 0, "after": 14},
            {
                "text": "It goes stale. The agent still treats it as the truth, and implements the stale parts with confidence.",
                "size": 18,
                "color": MUTED,
            },
        ],
        align=PP_ALIGN.LEFT,
        anchor="ctr",
    )
    shape_text(
        right,
        [
            {"text": "Kept current", "size": 14, "bold": True, "color": ACCENT_SOFT, "before": 0, "after": 8},
            {"text": "A short plan", "size": 28, "bold": True, "color": HEADER_INK, "before": 0, "after": 14},
            {
                "text": "One change, with acceptance, updated when the work moves. This is the planning that got more valuable.",
                "size": 18,
                "color": HEADER_INK,
            },
        ],
        align=PP_ALIGN.LEFT,
        anchor="ctr",
    )


def landscape_card(slide, x, name, stars, star_label, started, site, install, blurb):
    rect(slide, x, 1.85, 2.9, 3.85, CARD, line=LINE)
    rect(slide, x, 1.85, 2.9, 0.07, ACCENT)
    add_text(
        slide,
        [
            {"text": name, "size": 20, "bold": True, "after": 8},
            {"text": stars, "size": 26, "bold": True, "color": ACCENT, "after": 0},
            {"text": star_label, "size": 12, "color": MUTED, "after": 10},
            {"text": started, "size": 14, "after": 4},
            {"text": site, "size": 13, "color": MUTED, "after": 12},
            {"text": install, "size": 14, "bold": True, "after": 8},
            {"text": blurb, "size": 14, "color": MUTED},
        ],
        x + 0.16,
        2.08,
        2.58,
        3.8,
    )


def slide_landscape(prs):
    slide = new_slide(
        prs,
        3,
        """
        Read across the cards. These star counts and start dates came from the GitHub API today, 6 October 2026. They are context, not a ranking. Spec Kit and the example harness, Pi, are large. BMAD is large and a few months older. Squad is younger, smaller, and built for GitHub Copilot. The fourth column has nothing to install. The docs homepage for Spec Kit still says 130K stars and 38 integrations; the API count is 140,313, and the integrations table lists 40 named agents plus a generic one. Use the table and the API if anyone checks.
        """,
    )
    kicker_title(slide, "Landscape", "Four approaches")
    gap = 0.18
    width = 2.9
    x0 = 0.62
    landscape_card(
        slide,
        x0,
        "Spec Kit",
        "140,313",
        "GitHub stars",
        "Started 21 Aug 2025",
        "github.github.com/spec-kit",
        "Python 3.11+ and uv",
        "A spec through plan, tasks, implementation, and a check.",
    )
    landscape_card(
        slide,
        x0 + (width + gap),
        "BMAD",
        "53,827",
        "GitHub stars",
        "Started 13 Apr 2025",
        "docs.bmad-method.org",
        "Node, npm, and uv",
        "Planning depth that is supposed to size itself to the change.",
    )
    landscape_card(
        slide,
        x0 + 2 * (width + gap),
        "Squad",
        "3,251",
        "GitHub stars",
        "Started 6 Feb 2026",
        "bradygaster.github.io/squad",
        "Copilot, plus Node or a standalone install",
        "A roster of specialists that lives in the repository.",
    )
    landscape_card(
        slide,
        x0 + 3 * (width + gap),
        "Three files",
        "3 files",
        "Context, Plan, Roadmap",
        "No project installer",
        "Any harness that reads markdown",
        "Pi is the example harness",
        "112,833 stars. 15+ providers. Started 9 Aug 2025.",
    )


def slide_speckit(prs):
    slide = new_slide(
        prs,
        4,
        """
        Walk the loop. Constitution once, and only if those principles are already true. Then, per feature: specify, plan, tasks, implement, converge. Repeat implement and converge until it reports converged. Clarify, checklist, and analyze are extra quality gates. Install is Python 3.11+, uv, and an agent. The integrations reference lists 40 named agents plus generic; the marketing page still says 38. Brownfield, from their existing-project guide: you do not recreate the system as specifications. Do not make "document the entire existing system" the first feature. The new spec defines the change, and the codebase stays the implementation context. Started 21 August 2025. Release v1.1.0 on 2 October 2026. Pushed 5 October.
        """,
    )
    kicker_title(slide, "Spec Kit", "Specify the change, then build it")
    flow(slide, ["Specify", "Plan", "Tasks", "Implement", "Converge"], 0.62, 1.5, 12.1, h=0.5, size=15)
    add_text(
        slide,
        [
            {
                "text": "Constitution once per project. Only if those principles are already true.",
                "size": 15,
                "italic": True,
                "color": MUTED,
            }
        ],
        0.62,
        2.12,
        12.1,
        0.34,
    )
    bullets(
        slide,
        [
            "Python 3.11+, uv, and one of 40 named agents. A generic integration covers the rest.",
            "Clarify, checklist, and analyze are extra gates. The core loop stands without them.",
            "On an existing codebase, the spec is the next change. Leave the rest of the system in the code.",
        ],
        0.62,
        2.7,
        12.1,
        3.4,
        size=20,
    )
    meta_line(slide, "Started 21 Aug 2025    ·    140,313 stars    ·    v1.1.0 on 2 Oct 2026    ·    MIT")


def slide_bmad(prs):
    slide = new_slide(
        prs,
        5,
        """
        Use their own picture. A vague notion starts at clarify, a big clear idea at plan, a small change at build and verify, and learn-and-adjust loops back to plan. The README says small changes go straight to build. The heavier PRD, UX, and architecture path still exists for a change that is actually large. Personas and party mode are extra machinery: the default party is one model voicing every persona, and their docs say that single mind can quietly converge. Install is Node, npm, Git, and uv, via the skills CLI, or a Claude or Codex plugin marketplace. Started 13 April 2025. Latest release v6.12.1 on 4 October 2026. Pushed today.
        """,
    )
    kicker_title(slide, "BMAD", "Size the planning to the change")
    flow(slide, ["Clarify", "Plan", "Build and verify"], 0.62, 1.5, 12.1, h=0.5, size=16)
    add_text(
        slide,
        [
            {
                "text": "Vague notion, big idea, or small change. Learn and adjust loops back to plan.",
                "size": 15,
                "italic": True,
                "color": MUTED,
            }
        ],
        0.62,
        2.12,
        12.1,
        0.34,
    )
    bullets(
        slide,
        [
            "Node, npm, and uv, plus a tool that loads skills. Claude and Codex have plugin marketplaces too.",
            "Small changes go straight to build. Open the PRD and architecture path when the change is actually large.",
            "Personas and party mode sit on top of the loop. One model voicing five personas shares a mind.",
        ],
        0.62,
        2.7,
        12.1,
        3.4,
        size=20,
    )
    meta_line(slide, "Started 13 Apr 2025    ·    53,827 stars    ·    v6.12.1 on 4 Oct 2026    ·    MIT")


def slide_squad(prs):
    slide = new_slide(
        prs,
        6,
        """
        Squad's README says it is not a chatbot wearing hats: each member runs in its own context, reads its own knowledge, and writes back what it learned. That separate context is real. The charter is identity, expertise, and voice, compiled into the system prompt at spawn time. A tester is that prompt plus a routing rule. Their team docs include a tester when tests are detected. Day to day is copilot --agent squad. npm wants Node 22.5 or newer; Homebrew, WinGet, and the standalone installer are the other paths, and the standalone build vendors Node. Started 6 February 2026. Release v1.0.1 on 4 October 2026. Pushed today. 3,251 stars.
        """,
    )
    kicker_title(slide, "Squad", "A Copilot roster in the repo")
    flow(slide, ["Discover", "Propose a roster", "Spawn a role", "Human reviews"], 0.62, 1.5, 12.1, h=0.5, size=15)
    add_text(
        slide,
        [
            {
                "text": "The roster lives in .squad/. Charters, decisions, and history are files.",
                "size": 15,
                "italic": True,
                "color": MUTED,
            }
        ],
        0.62,
        2.12,
        12.1,
        0.34,
    )
    bullets(
        slide,
        [
            "Each member has its own context. The charter is compiled into the system prompt.",
            "A tester role is that prompt. Separate context helps. The job title leaves the model as it was.",
            "GitHub Copilot, plus their CLI. npm wants Node 22.5+. The standalone install vendors Node.",
        ],
        0.62,
        2.7,
        12.1,
        3.4,
        size=20,
    )
    meta_line(slide, "Started 6 Feb 2026    ·    3,251 stars    ·    v1.0.1 on 4 Oct 2026    ·    MIT")


def slide_files(prs):
    slide = new_slide(
        prs,
        7,
        """
        This is the whole method. Context is durable facts only: stack, conventions, domain terms, constraints, kept short. Plan is the current change: what, why, acceptance, approach, out of scope. The input is the requirement from the BA or product team, not a rediscovery. Roadmap is phases and items with status, updated when an item completes. That update is the agile record. Pi is an example harness, not a fourth method: minimal, 15 or more providers, and it ships without sub-agents or plan mode so you can keep your own workflow. Cursor or Copilot can read the same three files. No project installer.
        """,
    )
    kicker_title(slide, "Three files", "Context, plan, roadmap")
    flow(slide, ["Context.md", "Plan.md", "Roadmap.md"], 0.62, 1.5, 12.1, h=0.52, size=16)
    add_text(
        slide,
        [
            {
                "text": "Mark the item done. Fold anything still true next month back into context.",
                "size": 15,
                "italic": True,
                "color": MUTED,
            }
        ],
        0.62,
        2.16,
        12.1,
        0.34,
    )
    columns = [
        ("Context.md", "Durable facts only. Stack, conventions, domain terms, constraints. Kept short."),
        ("Plan.md", "The current change. What, why, acceptance, approach, and out of scope."),
        ("Roadmap.md", "Phases and items with status. Updated when an item is actually done."),
    ]
    for index, (title, body) in enumerate(columns):
        x = 0.62 + index * 4.1
        card = rect(slide, x, 2.7, 3.9, 2.45, CARD, line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        shape_text(
            card,
            [
                {"text": title, "size": 22, "bold": True, "color": ACCENT, "after": 12},
                {"text": body, "size": 16},
            ],
            anchor="ctr",
        )


def slide_fits(prs):
    slide = new_slide(
        prs,
        8,
        """
        Talk this as conditions, then leave the matrix up. Greenfield: you are writing decisions down as you make them. Spec Kit starts at specify. BMAD plans a big idea and builds a small change. Squad asks what you are building and proposes a team. The three files start as context you add when a decision sticks. Brownfield: Spec Kit's guide says the spec is the change, not a retroactive specification of every behaviour. BMAD says the knowledge is already in the source, and feeding the agent a prose replay creates contradictions and bloats the window. The three files verify a short context from the code and plan only the change. Squad scans the repo to cast roles, which answers a different question. Best used when: shared commands, optional depth, an existing Copilot shop, or the discipline to keep three files true.
        """,
    )
    kicker_title(slide, "Conditions", "Greenfield and brownfield")
    left = rect(slide, 0.62, 1.5, 6.0, 2.7, CARD, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    right = rect(slide, 6.82, 1.5, 5.9, 2.7, CARD, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    shape_text(
        left,
        [
            {"text": "Greenfield", "size": 14, "bold": True, "color": ACCENT, "after": 8},
            {"text": "Write decisions down as you make them.", "size": 20, "bold": True, "after": 10},
            {"text": "Spec Kit starts at specify. BMAD plans a big idea and builds a small change. Squad proposes a team from the description. The three files grow as decisions stick.", "size": 16},
        ],
        anchor="t",
    )
    shape_text(
        right,
        [
            {"text": "Brownfield", "size": 14, "bold": True, "color": ACCENT, "after": 8},
            {"text": "The code already holds the system.", "size": 20, "bold": True, "after": 10},
            {"text": "Spec Kit specs the next slice. BMAD warns against restating what the agent can read. The three files verify a short context and plan the change.", "size": 16},
        ],
        anchor="t",
    )
    add_text(
        slide,
        [
            {
                "text": "Shared commands → Spec Kit.     Optional depth → BMAD.     Already on Copilot → Squad.     You will keep three files true → stay with the files.",
                "size": 15,
                "color": MUTED,
            }
        ],
        0.62,
        4.45,
        12.1,
        0.9,
    )


def slide_matrix(prs):
    layout = next(l for l in prs.slide_layouts if l.name == "Blank")
    slide = prs.slides.add_slide(layout)
    paint(slide, PAPER)
    notes(
        slide,
        """
        Leave this up. Do not read every cell. Point at lock-in and install, then at the brownfield row. Invite people to photograph it. There is no diagram behind the table. Vanilla from the brief is the Three files column, so the audience can see what that column refuses to add: no CLI, no roster, no skill names.
        """,
    )
    add_text(
        slide,
        [{"text": "Side by side", "size": 22, "bold": True}],
        0.28,
        0.12,
        8,
        0.38,
    )
    add_text(
        slide,
        [{"text": "9  /  20", "size": 12, "color": MUTED}],
        10.3,
        0.18,
        2.75,
        0.28,
        align=PP_ALIGN.RIGHT,
    )

    headers = ["", "Spec Kit", "BMAD", "Squad", "Three files"]
    rows = [
        ["Unit of work", "A feature", "A sized change", "A task for a role", "One change"],
        ["Artifacts", "spec, plan, tasks", "spec, or a PRD if large", "charter, decisions, history", "Context, Plan, Roadmap"],
        ["Moving parts", "CLI and skills", "Skills and agents", "Roster and Copilot", "Three markdown files"],
        ["In common", "Intent, plan, tasks, a check", "Intent, plan, build, a check", "A task, decisions, a review", "Intent, plan, tasks, a check"],
        ["Lock-in", ".specify/ on many agents", "Skill names on many tools", "GitHub Copilot", "Any markdown reader"],
        ["Install", "Python 3.11+ and uv", "Node, npm, and uv", "Copilot and Node 22.5+, or standalone", "Nothing"],
        ["Greenfield", "Specify, then plan", "Plan the big idea", "Describe it, get a team", "Write context as you decide"],
        ["Brownfield", "The next slice only", "Do not restate the code", "Scan the repo, propose roles", "Verify context from the code"],
        ["Best used when", "Shared commands and a paper trail", "Depth you can leave off", "You already live in Copilot", "You will keep three files true"],
    ]

    label_w = 1.62
    left = 0.22
    top = 0.58
    table_w = W - 0.44
    col_w = (table_w - label_w) / 4
    header_h = 0.42
    row_h = (H - top - 0.16 - header_h) / len(rows)

    def cell(text, x, y, w, h, fill, color, size, bold, align):
        box = rect(slide, x, y, w, h, fill)
        shape_text(
            box,
            [{"text": text, "size": size, "bold": bold, "color": color}],
            align=align,
            anchor="ctr",
            size=size,
        )

    x = left
    for index, header in enumerate(headers):
        width = label_w if index == 0 else col_w
        cell(header, x, top, width, header_h, ACCENT, HEADER_INK, 13, True, PP_ALIGN.LEFT if index == 0 else PP_ALIGN.CENTER)
        x += width

    for r_index, row in enumerate(rows):
        y = top + header_h + r_index * row_h
        band = BAND if r_index % 2 == 0 else WHITE
        x = left
        for c_index, text in enumerate(row):
            width = label_w if c_index == 0 else col_w
            fill = LABEL if c_index == 0 else band
            cell(
                text,
                x,
                y,
                width,
                row_h,
                fill,
                INK,
                12 if c_index else 12,
                c_index == 0,
                PP_ALIGN.LEFT,
            )
            x += width


def slide_kernel(prs):
    slide = new_slide(
        prs,
        10,
        """
        The overlap is the method. The rest is packaging. Read down one column if you need to, then land on the four words at the bottom: written intent, a plan, tasks, a check. Squad's version of the check is a reviewer protocol, including a lockout so the author of the change cannot be the one who approves it. That separate pass is a real idea. The job titles are not the idea. Spec Kit's converge step is the same idea with less costume: repeat until the implementation matches the spec. BMAD's build-and-verify is the same idea, and the five-field spec they added — why, capabilities, constraints, non-goals, success — is a short plan.
        """,
    )
    kicker_title(slide, "In common", "They already agree")
    columns = [
        ("Spec Kit", ["Specify", "Plan", "Tasks", "Converge"]),
        ("BMAD", ["Spec", "Plan, if it is large", "Build", "Verify"]),
        ("Squad", ["The request", "Decisions file", "Routed work", "A review"]),
        ("Three files", ["What and why", "Approach and bounds", "The task list", "Acceptance"]),
    ]
    labels = ["Intent", "Plan", "Tasks", "Check"]
    for index, (name, steps) in enumerate(columns):
        x = 0.62 + index * 3.15
        card = rect(slide, x, 1.48, 3.0, 3.2, CARD, line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        paragraphs = [{"text": name, "size": 18, "bold": True, "color": ACCENT, "after": 8}]
        for step, label in zip(steps, labels):
            paragraphs.append({"text": label.upper(), "size": 11, "bold": True, "color": MUTED, "before": 6})
            paragraphs.append({"text": step, "size": 15, "after": 2})
        shape_text(card, paragraphs, anchor="ctr")
    add_text(
        slide,
        [
            {
                "text": "Written intent.  A plan.  Tasks.  A check.",
                "size": 18,
                "bold": True,
                "color": ACCENT,
            }
        ],
        0.62,
        4.9,
        12.1,
        0.4,
    )


def rent_row(slide, y, number, title, body):
    bubble = rect(slide, 0.62, y, 0.46, 0.46, ACCENT, shape=MSO_SHAPE.OVAL)
    shape_text(
        bubble,
        [{"text": number, "size": 16, "bold": True, "color": HEADER_INK}],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    add_text(
        slide,
        [
            {"text": title, "size": 20, "bold": True, "after": 2},
            {"text": body, "size": 16, "color": MUTED},
        ],
        1.28,
        y - 0.06,
        11.4,
        1.15,
    )


def slide_rent(prs):
    slide = new_slide(
        prs,
        11,
        """
        Three things that do not earn a place in the default path. First, a tester persona. Squad compiles the charter into the prompt. BMAD measured a cousin of this idea: the v6.11.0 notes, 10 August 2026, say an A/B on Claude and Codex found the cynical, jaded reviewer framing made no difference to residual-bug hit rate, while requiring at least ten concrete findings and asking what is missing did. Quote that if challenged. Second, a standing model per role. Squad's parallel-work docs route by task and by charter, and say cost wins when in doubt. Their switching-models scenario puts the tester on the fast tier and says the tester does not need the premium model to write tests. That catalog is a maintenance job. Keep one strong model and a separate review pass. Third, phase gates for a change whose requirements are already known. Fresh context plus an explicit check is the part worth keeping.
        """,
    )
    kicker_title(slide, "Cost without a return", "What does not pay rent")
    rent_row(
        slide,
        1.6,
        "1",
        "A tester charter",
        "BMAD’s v6.11 notes: a cynical reviewer persona did not change the residual-bug hit rate. Concrete findings did.",
    )
    rent_row(
        slide,
        3.15,
        "2",
        "A model pinned to a role",
        "Squad maps roles onto price tiers, and a charter can pin a model. The map rots as models ship. Keep one strong model.",
    )
    rent_row(
        slide,
        4.7,
        "3",
        "A phase gate for a known change",
        "If the requirement is already bounded, more ceremony writes more pages before the same check.",
    )


def slide_ceremony(prs):
    slide = new_slide(
        prs,
        12,
        """
        Agile, in this talk, means one bounded change in flight and then an update to the record. A phase gate is specifying the whole product before anyone sees code. Do not say waterfall died. People in this room still ship phase gates, and the line is too easy to swat. Say the heavy path recreates phase gates, and that each project now ships a short path because the long one is too much. Evidence: Spec Kit is per feature, and bug fixing does not require the feature workflow first. BMAD says small changes go straight to build. You can still turn the heavy setting on. The tool is not a waterfall. The heavy setting feels like one.
        """,
    )
    kicker_title(slide, "Pace", "One change in flight")
    left = rect(slide, 0.62, 1.6, 6.0, 4.55, ACCENT, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    right = rect(slide, 6.82, 1.6, 5.9, 4.55, CARD, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    shape_text(
        left,
        [
            {"text": "The short path", "size": 14, "bold": True, "color": ACCENT_SOFT, "after": 10},
            {"text": "One bounded change, then update the record.", "size": 24, "bold": True, "color": HEADER_INK, "after": 12},
            {"text": "Spec Kit runs the loop per feature. BMAD sends a small change straight to build. Squad waits for a person to accept the work.", "size": 16, "color": HEADER_INK},
        ],
        anchor="ctr",
    )
    shape_text(
        right,
        [
            {"text": "The heavy setting", "size": 14, "bold": True, "color": ACCENT, "after": 10},
            {"text": "The whole product, specified before any code.", "size": 24, "bold": True, "after": 12},
            {"text": "That setting recreates a phase gate. Each project also ships a short path, because the long one is too much.", "size": 16, "color": MUTED},
        ],
        anchor="ctr",
    )


def slide_lockin(prs):
    slide = new_slide(
        prs,
        13,
        """
        Split lock-in into three, or the question is mush. Model lock-in is low for Spec Kit and BMAD and high for Squad, because Squad is a Copilot team. Framework lock-in is real for all three: the .specify directory, BMAD skill names, and the .squad directory. Markdown lock-in is the acceptable kind, because the files move with you. Upkeep: abandonment is the wrong charge. Spec Kit was pushed on 5 October, and BMAD, Squad, and Pi were pushed on 6 October. BMAD v6.11, published 10 August 2026, renamed bmad-quick-dev to bmad-build and left shims until a v7 cut. The latest tag on top of that is v6.12.1, 4 October. They are maintained. The cost is keeping up with them. Licenses are MIT. BMAD's GitHub API license field says NOASSERTION because the LICENSE file adds a contributor paragraph; the file itself is MIT.
        """,
    )
    kicker_title(slide, "Switching cost", "Lock-in and upkeep")
    rows = [
        ("Model", "Low for Spec Kit and BMAD. High for Squad: the product is a Copilot team."),
        ("Framework", ".specify/, BMAD skill names, and .squad/ stay behind when you leave."),
        ("Markdown", "Context, plan, and roadmap move to the next harness. That lock-in is cheap."),
    ]
    for index, (name, body) in enumerate(rows):
        y = 1.55 + index * 1.25
        rect(slide, 0.62, y, 1.7, 1.05, ACCENT, radius=0.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(
            slide,
            [{"text": name, "size": 16, "bold": True, "color": HEADER_INK}],
            0.62,
            y,
            1.7,
            1.05,
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        add_text(
            slide,
            [{"text": body, "size": 18}],
            2.55,
            y,
            10.1,
            1.05,
            anchor="ctr",
        )
    add_text(
        slide,
        [
            {
                "text": "All four repos were pushed on 5 or 6 Oct 2026. BMAD renamed its build command in August and is shimming skills toward a v7 cut.",
                "size": 14,
                "color": MUTED,
            }
        ],
        0.62,
        5.4,
        12.1,
        0.7,
    )


def slide_recommend(prs):
    slide = new_slide(
        prs,
        14,
        """
        Walk the loop slowly. The BA or product requirement is the input. It goes into Plan.md: what, why, acceptance, approach, out of scope. You build that change. You check it against acceptance. You mark the roadmap item done. Anything that will still be true next month is folded into Context, and Context stays short. This is the same loop the frameworks are collapsing toward. It works on Pi, Cursor, Copilot, or anything else that will read the files. There is no project installer and no roster to maintain.
        """,
    )
    kicker_title(slide, "Recommendation", "How a requirement moves")
    flow(
        slide,
        ["BA requirement", "Plan.md", "Build", "Check", "Roadmap"],
        0.5,
        1.5,
        12.3,
        h=0.52,
        size=14,
    )
    cards = [
        ("Context.md", "Durable facts. Correct it when it is wrong. Do not let it become a second codebase."),
        ("Plan.md", "One change. The requirement is the input. Acceptance is the check."),
        ("Roadmap.md", "The agile record. Update the status when the item is done. It is not a freeze."),
    ]
    for index, (title, body) in enumerate(cards):
        x = 0.62 + index * 4.1
        card = rect(slide, x, 2.3, 3.9, 2.45, CARD, line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        shape_text(
            card,
            [
                {"text": title, "size": 22, "bold": True, "color": ACCENT, "after": 12},
                {"text": body, "size": 18},
            ],
            anchor="ctr",
        )


def slide_both(prs):
    slide = new_slide(
        prs,
        15,
        """
        The loop does not change between a new system and an old one. What changes is where Context comes from. Greenfield: write a decision into Context when it sticks. Brownfield: verify a short context from the code, and plan only the change. Spec Kit: do not document the entire existing system unless that inventory is the deliverable. The spec defines the change you intend, not every current behaviour. BMAD's existing-codebase guide: most of the knowledge is already in the source, and textual descriptions of things the agent can already read create contradiction, ambiguity, and context-window bloat. Their older document-the-project workflow is deprecated. Both codebases still take the requirement from the BA or product team into Plan.md.
        """,
    )
    kicker_title(slide, "One method", "Same loop, both codebases")
    left = rect(slide, 0.62, 1.55, 6.0, 4.7, CARD, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    right = rect(slide, 6.82, 1.55, 5.9, 4.7, ACCENT, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    shape_text(
        left,
        [
            {"text": "Greenfield", "size": 14, "bold": True, "color": ACCENT, "after": 8},
            {"text": "Context is the decisions.", "size": 24, "bold": True, "after": 12},
            {"text": "Write a fact when it sticks. Plan the change you were given. Mark the roadmap when it ships.", "size": 18},
        ],
        anchor="ctr",
    )
    shape_text(
        right,
        [
            {"text": "Brownfield", "size": 14, "bold": True, "color": ACCENT_SOFT, "after": 8},
            {"text": "Context is verified from the code.", "size": 24, "bold": True, "color": HEADER_INK, "after": 12},
            {"text": "A short checked context. A plan for the next change. Restating the system creates a second, staler source.", "size": 18, "color": HEADER_INK},
        ],
        anchor="ctr",
    )


def slide_unknown(prs):
    slide = new_slide(
        prs,
        16,
        """
        The caveat, said plainly. If you do not know the requirement, do not write a plan that pretends you do. The agent will implement the fiction. Discover until the change is bounded: who it is for, what done looks like, what is out of scope. Then write Plan.md. This case is uncommon when a BA or product team is feeding the work. It is fatal when it is skipped. A framework does not rescue it. BMAD's own spec skill says the spec writes the contract and does not help you figure out what you want; input too thin to use is sent back. Spec Kit's idea-assessment path can end in a documented stop. That stop is a legitimate result. More pages of guesses are still guesses.
        """,
    )
    kicker_title(slide, "Caveat", "If the requirement is unknown")
    add_text(
        slide,
        [
            {
                "text": "Discover until the change is bounded.\nThen write the plan.",
                "size": 28,
                "bold": True,
            }
        ],
        0.62,
        1.5,
        12.1,
        1.3,
    )
    flow(slide, ["Ask", "Bound the change", "Plan.md"], 0.62, 3.1, 12.1, h=0.56, size=18)
    add_text(
        slide,
        [
            {
                "text": "Uncommon when a product team is feeding the work. Fatal when it is skipped. Every framework fails here the same way: a thick spec made of guesses is still a guess.",
                "size": 18,
                "color": MUTED,
            }
        ],
        0.62,
        4.0,
        12.1,
        1.4,
    )


def slide_package(prs):
    slide = new_slide(
        prs,
        17,
        """
        The concession, so this is not a slogan. If a team will not keep three files honest, and they want shared commands plus a paper trail, use Spec Kit's core loop: specify, plan, tasks, implement, converge. Leave the extra gates off unless a particular change is ambiguous. Leave constitution alone unless the principles are ones you already mean. BMAD's five-field spec — why, capabilities, constraints, non-goals, success signal — is a Plan.md with a skill wrapped around it. Use that shape. Do not turn the full PRD and architecture path on for a change you already understand. Do not adopt Squad in order to grow a better tester. Take Squad only if you already live in Copilot and you want the separate contexts, with your eyes open about the roster and the model map.
        """,
    )
    kicker_title(slide, "Concession", "If you still want a package")
    left = rect(slide, 0.62, 1.55, 6.0, 4.7, ACCENT, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    right = rect(slide, 6.82, 1.55, 5.9, 4.7, CARD, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    shape_text(
        left,
        [
            {"text": "Use", "size": 14, "bold": True, "color": ACCENT_SOFT, "after": 8},
            {"text": "Spec Kit’s core loop.", "size": 26, "bold": True, "color": HEADER_INK, "after": 12},
            {"text": "Specify, plan, tasks, implement, converge. One feature at a time. Their five-step cousin in BMAD is the same short spec: why, capabilities, constraints, non-goals, success.", "size": 16, "color": HEADER_INK},
        ],
        anchor="ctr",
    )
    shape_text(
        right,
        [
            {"text": "Leave off", "size": 14, "bold": True, "color": ACCENT, "after": 8},
            {"text": "The extra machinery.", "size": 26, "bold": True, "after": 12},
            {"text": "Gates you do not need. A constitution you do not mean. A roster adopted to grow a better tester. BMAD’s full path for a change you already understand.", "size": 16, "color": MUTED},
        ],
        anchor="ctr",
    )


def slide_close(prs):
    slide = new_slide(
        prs,
        18,
        """
        Close on the two lines. Keep the plan current. Do not adopt a department of prompts. Pause. If there is time, take questions. The appendix has the URLs and the retrieval date. Do not walk the appendix unless someone asks where a number came from. If someone asks "so you are saying never use them," go back to the concession: Spec Kit's core loop, extra gates off, when a team needs the shared commands.
        """,
    )
    add_text(
        slide,
        [{"text": "CLOSE", "size": 13, "bold": True, "color": ACCENT}],
        0.7,
        1.45,
        12,
        0.3,
    )
    add_text(
        slide,
        [{"text": "Keep the plan current.", "size": 40, "bold": True}],
        0.7,
        1.9,
        12,
        0.7,
    )
    add_text(
        slide,
        [{"text": "Do not adopt a department of prompts.", "size": 28, "color": MUTED}],
        0.7,
        2.75,
        12,
        0.55,
    )
    flow(slide, ["Context.md", "Plan.md", "Roadmap.md"], 0.7, 4.0, 11.9, h=0.6, size=18)
    add_text(
        slide,
        [
            {
                "text": "The files are the method. The harness is the one you already have.",
                "size": 18,
                "color": MUTED,
            }
        ],
        0.7,
        4.9,
        11.5,
        0.4,
    )


def source_slide(prs, number, kicker, title, columns, spoken):
    slide = new_slide(prs, number, spoken)
    kicker_title(slide, kicker, title)
    for index, blocks in enumerate(columns):
        add_text(slide, blocks, 0.62 + index * 6.2, 1.45, 5.95, 5.4)
    return slide


def src(text, *, bold=False, gap=3, size=13, color=INK):
    return {"text": text, "size": size, "bold": bold, "before": gap, "color": color}


def slide_sources_a(prs):
    source_slide(
        prs,
        19,
        "Appendix",
        "Sources, retrieved 6 Oct 2026",
        [
            [
                src("Spec Kit", bold=True, gap=0, size=16, color=ACCENT),
                src("github.com/github/spec-kit — API 6 Oct 2026"),
                src("140,313 stars. Created 21 Aug 2025. Pushed 5 Oct 2026. MIT."),
                src("Release v1.1.0 published 2 Oct 2026."),
                src("README: Python 3.11+, uv; constitution once; specify → plan → tasks → implement → converge."),
                src("Integrations table: 40 named agents, plus generic. Docs homepage still says 38 and 130K+ stars."),
                src("Existing-project guide: do not recreate the system as specs; spec the next change."),
                src("BMAD", bold=True, gap=12, size=16, color=ACCENT),
                src("github.com/bmad-code-org/BMAD-METHOD — API 6 Oct 2026"),
                src("53,827 stars. Created 13 Apr 2025. Pushed 6 Oct 2026."),
                src("LICENSE file is MIT. API SPDX field is NOASSERTION."),
                src("Release v6.12.1 on 4 Oct 2026. v6.11.0 on 10 Aug 2026 renamed quick-dev to build; shims until v7; persona A/B."),
            ],
            [
                src("BMAD, continued", bold=True, gap=0, size=16, color=ACCENT),
                src("README: skills CLI, Claude and Codex marketplaces, uv. Small changes go straight to build."),
                src("Delivery loop: clarify, plan, build and verify; learn loops back."),
                src("docs.bmad-method.org — existing codebase, party mode, define a specification."),
                src("Party mode: one model voicing five personas can quietly converge."),
                src("Spec kernel: Why, Capabilities, Constraints, Non-goals, Success signal."),
                src("Existing codebase: restating readable code bloats the window. document-project is deprecated."),
                src("Full URLs are in sources.md, with the same retrieval date.", gap=14, color=MUTED),
            ],
        ],
        "Not spoken unless someone asks where a Spec Kit or BMAD figure came from. Point them at sources.md as well.",
    )


def slide_sources_b(prs):
    source_slide(
        prs,
        20,
        "Appendix",
        "Sources, continued",
        [
            [
                src("Squad", bold=True, gap=0, size=16, color=ACCENT),
                src("github.com/bradygaster/squad — API 6 Oct 2026"),
                src("3,251 stars. Created 6 Feb 2026. Pushed 6 Oct 2026. MIT."),
                src("Release v1.0.1 published 4 Oct 2026. Homepage bradygaster.github.io/squad."),
                src("README: Copilot team; not a chatbot wearing hats; each member has its own context."),
                src("npm install wants Node.js 22.5+. Standalone install vendors Node. copilot --agent squad."),
                src("Your Team: charter.md is identity, expertise, and voice, compiled into the prompt."),
                src("Parallel Work and Models: per-role and per-task model routing; cost wins when in doubt."),
                src("Switching Models scenario: tester on the fast tier."),
            ],
            [
                src("Pi, the example harness", bold=True, gap=0, size=16, color=ACCENT),
                src("github.com/earendil-works/pi — API 6 Oct 2026"),
                src("112,833 stars. Created 9 Aug 2025. Pushed 6 Oct 2026. MIT."),
                src("Release v1.0.4 published 5 Oct 2026. Site pi.dev."),
                src("pi.dev: minimal harness; 15+ providers; skips sub-agents and plan mode in the core."),
                src("Install: curl https://pi.dev/install.sh, or npm i -g @earendil-works/pi-coding-agent."),
                src("Pi is not a fourth method. It is one harness that will read the three files.", gap=10),
                src("Star counts and timestamps are from the GitHub REST API on 6 Oct 2026. Integration count is a manual count of the Spec Kit integrations table the same day.", gap=12, color=MUTED),
            ],
        ],
        "Not spoken unless someone asks about Squad, Pi, or how the star counts were taken.",
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Do the new AI frameworks earn their keep?"
    prs.core_properties.subject = "AI-driven development meetup, 6 October 2026"
    prs.core_properties.category = "Meetup slides"
    builders = [
        slide_title,
        slide_shift,
        slide_landscape,
        slide_speckit,
        slide_bmad,
        slide_squad,
        slide_files,
        slide_fits,
        slide_matrix,
        slide_kernel,
        slide_rent,
        slide_ceremony,
        slide_lockin,
        slide_recommend,
        slide_both,
        slide_unknown,
        slide_package,
        slide_close,
        slide_sources_a,
        slide_sources_b,
    ]
    for builder in builders:
        builder(prs)
    if len(prs.slides) != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, got {len(prs.slides)}")
    # Matrix is slide 9 and must not carry a decorative bar. Confirm no picture parts.
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type is not None and shape.shape_type == 13:
                raise SystemExit("picture shape found; this deck uses native shapes only")
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
